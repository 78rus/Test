# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

# Each entry: (number, question_text, [(option_text, is_correct), ...])
questions = [
    (6, "The main goal of epidemiology is to ___________ diseases.",
        [("prevent and control", True), ("treat only", False), ("diagnose only", False)]),
    (7, "Rospotrebnadzor is responsible for ___________ in Russia.",
        [("space exploration", False), ("consumer protection", True), ("transportation", False)]),
    (8, "The most effective way to prevent healthcare-associated infections is ___________.",
        [("wearing gloves only", False), ("hand hygiene", True), ("isolation", False)]),
    (9, "A patient with a weakened immune system is most vulnerable to ___________.",
        [("allergies", False), ("genetic disorders", False), ("opportunistic infections", True)]),
    (10, "The nurse ___________ the patient's temperature now.",
        [("take", False), ("takes", False), ("is taking", True)]),
    (11, "A ___________ is used to monitor heart activity.",
        [("MRI", False), ("ECG", True), ("CT scan", False)]),
    (12, "The first line of defense against infectious diseases is ___________.",
        [("treatment", False), ("surgery", False), ("hygiene", True)]),
    (13, "A ___________ is a solid dose form inserted into the rectum.",
        [("tablet", False), ("capsule", False), ("suppository", True)]),
    (14, "Occupational hygiene focuses on ___________.",
        [("workplace safety and health", True), ("food safety", False), ("school hygiene", False)]),
    (15, "The most common cause of foodborne illness is ___________.",
        [("chemicals", False), ("bacteria", True), ("radiation", False)]),
    (16, "___________ is used to examine soft tissues in the body.",
        [("MRI", True), ("ECG", False), ("X-ray", False)]),
    (17, "The main purpose of consumer protection is to ___________.",
        [("increase prices", False), ("protect buyers from unfair practices", True), ("promote advertising", False)]),
    (18, "___________ is the practice of keeping oneself clean to prevent disease.",
        [("Radiation", False), ("Medication", False), ("Hygiene", True)]),
    (19, "A ___________ is a liquid medication taken by mouth.",
        [("tablet", False), ("capsule", False), ("syrup", True)]),
    (20, "The spread of infectious diseases is monitored by ___________.",
        [("epidemic surveillance", True), ("food inspection", False), ("environmental protection", False)]),
    (21, "___________ is the process of destroying all microorganisms on a surface.",
        [("Cleaning", False), ("Sterilization", True), ("Disinfection", False)]),
    (22, "The most effective way to reduce occupational stress is ___________.",
        [("increasing workload", False), ("promoting work-life balance", True), ("reducing salary", False)]),
    (23, "___________ is used to examine the stomach lining.",
        [("Colonoscopy", False), ("MRI", False), ("Gastroscopy", True)]),
    (24, "A ___________ is applied to the skin.",
        [("tablet", False), ("injection", False), ("ointment", True)]),
    (25, "The main source of water pollution in cities is ___________.",
        [("parks", False), ("industrial waste", True), ("playgrounds", False)]),
    (26, "___________ is a type of scan that uses X-rays to create cross-sectional images.",
        [("MRI", False), ("CT", True), ("ECG", False)]),
    (27, "Poor personal hygiene can lead to ___________.",
        [("improved health", False), ("better skin", False), ("infectious diseases", True)]),
    (28, "The role of Rospotrebnadzor includes ___________.",
        [("building roads", False), ("consumer rights protection", True), ("teaching languages", False)]),
    (29, "___________ is used to treat asthma.",
        [("Tablet", False), ("Injection", False), ("Inhaler", True)]),
    (30, "The main types of radiation include ___________.",
        [("sound and light", False), ("ionizing and non-ionizing", True), ("hot and cold", False)]),
    (31, "___________ is a measure to reduce air pollution in cities.",
        [("Increasing traffic", False), ("Using public transport", True), ("Building more factories", False)]),
    (32, "___________ is a type of personal protective equipment for workers.",
        [("Apron", False), ("Respirator", True), ("Scarf", False)]),
    (33, "Community hygiene studies the impact of ___________ on health.",
        [("education", False), ("environment", True), ("economy", False)]),
    (34, "The most common infectious disease in schools is ___________.",
        [("diabetes", False), ("influenza", True), ("hypertension", False)]),
    (35, "___________ is used to deliver medication directly into the bloodstream.",
        [("Tablet", False), ("Injection", True), ("Suppository", False)]),
    (36, "The main goal of urban ecology is to ___________.",
        [("increase pollution", False), ("improve the quality of life", True), ("reduce green spaces", False)]),
    (37, "___________ is an example of a biological occupational hazard.",
        [("Noise", False), ("Radiation", False), ("Bacteria", True)]),
    (38, "The purpose of food hygiene is to ___________.",
        [("make food tastier", False), ("prevent foodborne diseases", True), ("increase shelf life only", False)]),
    (39, "___________ is a way to prevent the spread of germs in public places.",
        [("Sharing food", False), ("Regular handwashing", True), ("Ignoring coughs", False)]),
    (40, "A ___________ is used to measure blood pressure.",
        [("thermometer", False), ("sphygmomanometer", True), ("stethoscope", False)]),
    (41, "___________ is a common cause of noise pollution in cities.",
        [("Trees", False), ("Traffic", True), ("Parks", False)]),
    (42, "The main function of white blood cells is to ___________.",
        [("carry oxygen", False), ("fight infection", True), ("clot blood", False)]),
    (43, "___________ is the branch of medicine that studies the heart.",
        [("Neurology", False), ("Cardiology", True), ("Dermatology", False)]),
    (44, "The best way to prevent the spread of COVID-19 is ___________.",
        [("ignoring symptoms", False), ("vaccination and hand hygiene", True), ("sharing food", False)]),
    (45, "___________ is a type of alternative medicine.",
        [("Acupuncture", True), ("Antibiotics", False), ("Surgery", False)]),
    (46, "The main source of radiation exposure for the general public is ___________.",
        [("nuclear power plants", False), ("medical imaging", True), ("mobile phones", False)]),
    (47, "___________ is an example of a chronic disease.",
        [("Influenza", False), ("Diabetes", True), ("Food poisoning", False)]),
    (48, "The main purpose of epidemic surveillance is ___________.",
        [("early detection of outbreaks", True), ("treating patients", False), ("testing drugs", False)]),
    (49, "___________ is a physical occupational factor.",
        [("Noise", True), ("Stress", False), ("Boredom", False)]),
    (50, "The most effective way to promote hygiene in schools is ___________.",
        [("punishment", False), ("education and awareness", True), ("ignoring the issue", False)]),
    (51, "___________ is used to examine the colon.",
        [("Colonoscopy", True), ("Gastroscopy", False), ("MRI", False)]),
    (52, "The main function of platelets is to ___________.",
        [("carry oxygen", False), ("fight infection", False), ("help blood clot", True)]),
    (53, "___________ is a governmental body responsible for consumer protection in Russia.",
        [("Ministry of Finance", False), ("Rospotrebnadzor", True), ("Ministry of Sports", False)]),
    (54, "The most common symptom of food poisoning is ___________.",
        [("fever only", False), ("nausea and vomiting", True), ("headache", False)]),
    (55, "___________ is a way to reduce land pollution.",
        [("Recycling", True), ("Burning waste", False), ("Using more plastic", False)]),
    (56, "A ___________ is used to administer medication via inhalation.",
        [("tablet", False), ("injection", False), ("inhaler", True)]),
    (57, "The main challenge in public health today is ___________.",
        [("lack of hospitals", False), ("antibiotic resistance", True), ("too many doctors", False)]),
    (58, "___________ is an ergonomic factor in occupational hygiene.",
        [("Noise", False), ("Posture", True), ("Radiation", False)]),
    (59, "The main source of indoor air pollution is ___________.",
        [("plants", False), ("tobacco smoke", True), ("sunlight", False)]),
    (60, "___________ is a long-term health effect of radiation exposure.",
        [("Headache", False), ("Cancer", True), ("Skin rash", False)]),
    (61, "The main function of the heart is to ___________.",
        [("digest food", False), ("pump blood", True), ("filter waste", False)]),
    (62, "___________ is a way to ensure food safety at home.",
        [("Eating raw meat", False), ("Proper cooking and storage", True), ("Washing dishes occasionally", False)]),
    (63, "The most common mental illness in adults is ___________.",
        [("schizophrenia", False), ("depression", True), ("bipolar disorder", False)]),
    (64, "___________ is a type of radiation hygiene measure.",
        [("Shielding", True), ("Smoking", False), ("Exercising", False)]),
    (65, "The main cause of waterborne diseases is ___________.",
        [("noise", False), ("contaminated water", True), ("light pollution", False)]),
    (66, "___________ is an example of a personal hygiene practice.",
        [("Driving fast", False), ("Brushing teeth", True), ("Eating junk food", False)]),
    (67, "The main goal of occupational hygiene is to ___________.",
        [("prevent work-related illnesses", True), ("increase profits", False), ("reduce salaries", False)]),
    (68, "___________ is a type of community hygiene measure.",
        [("Building hospitals", False), ("Waste management", True), ("Teaching literature", False)]),
    (69, "The most common cause of death from infectious disease is ___________.",
        [("diabetes", False), ("pneumonia", True), ("cancer", False)]),
    (70, "___________ is a type of consumer fraud.",
        [("Discounts", False), ("False advertising", True), ("Free samples", False)]),
    (71, "The best way to prevent foodborne illness in a restaurant is ___________.",
        [("using fewer ingredients", False), ("proper food storage and hand hygiene", True), ("serving raw food", False)]),
    (72, "___________ is a psychological occupational stressor.",
        [("Noise", False), ("High workload", True), ("Heat", False)]),
    (73, "The main function of the kidneys is to ___________.",
        [("pump blood", False), ("digest food", False), ("filter waste from blood", True)]),
    (74, "___________ is a way to reduce noise pollution.",
        [("Playing loud music", False), ("Using sound barriers", True), ("Ignoring the problem", False)]),
    (75, "The most common cause of dental disease is ___________.",
        [("genetics only", False), ("poor oral hygiene", True), ("drinking water", False)]),
    (76, "___________ is a type of investigation used to diagnose bone fractures.",
        [("MRI", False), ("CT scan", False), ("X-ray", True)]),
    (77, "The main purpose of quarantine is to ___________.",
        [("prevent the spread of disease", True), ("treat patients", False), ("test vaccines", False)]),
    (78, "___________ is an example of a chemical hazard in the workplace.",
        [("Noise", False), ("Radiation", False), ("Toxic fumes", True)]),
    (79, "The most effective way to reduce air pollution is ___________.",
        [("burning more fuel", False), ("using renewable energy", True), ("cutting trees", False)]),
    (80, "___________ is a type of medication used to lower blood pressure.",
        [("Antibiotic", False), ("Antihypertensive", True), ("Analgesic", False)]),
    (81, "The main cause of occupational hearing loss is ___________.",
        [("music", False), ("stress", False), ("prolonged noise exposure", True)]),
    (82, "___________ is a branch of medicine that studies mental disorders.",
        [("Cardiology", False), ("Dermatology", False), ("Psychiatry", True)]),
    (83, "The best way to prevent the spread of infections in schools is ___________.",
        [("closing schools", False), ("educating students about hygiene", True), ("ignoring sick children", False)]),
    (84, "___________ is a type of alternative medicine.",
        [("Surgery", False), ("Herbal medicine", True), ("Chemotherapy", False)]),
    (85, "The main function of the liver is to ___________.",
        [("pump blood", False), ("detoxify the body", True), ("filter air", False)]),
    (86, "___________ is a measure to prevent radiation exposure.",
        [("Smoking", False), ("Using shielding", True), ("Exercising", False)]),
    (87, "The most common cause of respiratory infections is ___________.",
        [("poor diet", False), ("viruses and bacteria", True), ("lack of sleep", False)]),
    (88, "___________ is a type of consumer protection measure.",
        [("Increasing prices", False), ("Product labelling", True), ("Reducing quality", False)]),
    (89, "The main purpose of community hygiene is to ___________.",
        [("increase pollution", False), ("promote public health", True), ("reduce green spaces", False)]),
]

doc = Document()

# Base style
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.font.bold = False
rpr = style.element.get_or_add_rPr()
rfonts = rpr.get_or_add_rFonts()
rfonts.set(qn('w:ascii'), 'Times New Roman')
rfonts.set(qn('w:hAnsi'), 'Times New Roman')
rfonts.set(qn('w:cs'), 'Times New Roman')
pf = style.paragraph_format
pf.line_spacing = 1.0
pf.space_before = Pt(0)
pf.space_after = Pt(0)

from docx.enum.table import WD_ALIGN_VERTICAL

def set_cell_text(cell, text, align, vcenter=False):
    cell.text = ""
    if vcenter:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = False
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:ascii'), 'Times New Roman')
    rf.set(qn('w:hAnsi'), 'Times New Roman')
    rf.set(qn('w:cs'), 'Times New Roman')

def add_plain(text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = bold
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:ascii'), 'Times New Roman')
    rf.set(qn('w:hAnsi'), 'Times New Roman')
    rf.set(qn('w:cs'), 'Times New Roman')
    return p

CENTER = WD_ALIGN_PARAGRAPH.CENTER
LEFT = WD_ALIGN_PARAGRAPH.LEFT

from docx.shared import Cm

def set_col_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    # fixed layout
    layout = tblPr.find(qn('w:tblLayout'))
    if layout is None:
        layout = tblPr.makeelement(qn('w:tblLayout'), {})
        tblPr.append(layout)
    layout.set(qn('w:type'), 'fixed')
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w

COL_WIDTHS = [Cm(2.4), Cm(11.5), Cm(3.5)]

for num, qtext, options in questions:
    add_plain("%d. %s" % (num, qtext))

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    hdr = table.rows[0].cells
    set_cell_text(hdr[0], "Поле для выбора ответа", CENTER, vcenter=True)
    set_cell_text(hdr[1], "Варианты ответов", CENTER, vcenter=True)
    set_cell_text(hdr[2], "Поле для отметки правильного ответа (+)", CENTER)

    for opt_text, correct in options:
        row = table.add_row().cells
        set_cell_text(row[0], "", CENTER, vcenter=True)          # empty selection field, centered H+V
        set_cell_text(row[1], opt_text, CENTER, vcenter=True)    # option text, centered H+V
        set_cell_text(row[2], "+" if correct else "-", LEFT)     # +/- left aligned

    set_col_widths(table, COL_WIDTHS)

doc.save("quiz.docx")
print("saved quiz.docx with", len(questions), "questions")
