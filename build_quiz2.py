# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn

# (number, question_text, [(option_text, is_correct), ...])
questions = [
    (11, "\u201cWhere does it hurt?\u201d = He asked ___________.",
        [("where does it hurt", False), ("where did it hurt", False), ("where it hurts", True)]),
    (12, "Radiology is the branch of medicine that uses ___________________________.",
        [("surgical instruments", False), ("medications", False), ("imaging techniques to diagnose diseases", True), ("physical therapy", False)]),
    (13, "Cardiology is the branch of medicine that studies ___________________________.",
        [("the brain", False), ("the heart and blood vessels", True), ("the skin", False), ("the digestive system", False)]),
    (14, "Neurology deals with ___________________________.",
        [("bones and joints", False), ("the heart and blood vessels", False), ("disorders of the nervous system", True), ("kidney diseases", False)]),
    (15, "Ophthalmology specializes in ___________________________.",
        [("teeth and gums", False), ("eye diseases and vision", True), ("skin diseases", False), ("blood disorders", False)]),
    (16, "Haematology is the study of ___________________________.",
        [("bones and joints", False), ("the heart and blood vessels", False), ("the nervous system", False), ("blood and blood disorders", True)]),
    (17, "Gastroenterology investigates and treats ___________________________.",
        [("skin diseases", False), ("heart conditions", False), ("kidney diseases", False), ("disorders of the digestive system", True)]),
    (18, "A general practitioner (GP) ___________________________.",
        [("performs surgery", False), ("treats only children", False), ("provides primary care for all ages", True), ("specializes in skin diseases", False)]),
    (19, "Physiotherapy ___________________________.",
        [("treats diseases of the skin", False), ("deals with sick children", False), ("designs exercises for patients", True), ("performs operations on patients", False)]),
    (20, "A pulmonologist treats ___________________________.",
        [("heart diseases", False), ("lung and respiratory disorders", True), ("kidney diseases", False), ("skin diseases", False)]),
    (21, "An ECG is used to ___________________________.",
        [("examine the brain", False), ("monitor heart activity", True), ("detect bone fractures", False), ("examine the stomach", False)]),
    (22, "An MRI is used to ___________________________.",
        [("examine soft tissues", True), ("monitor heart activity", False), ("detect bone fractures", False), ("measure blood pressure", False)]),
    (23, "A CT scan is used to ___________________________.",
        [("examine soft tissues", False), ("monitor heart activity", False), ("create cross-sectional images", True), ("measure blood pressure", False)]),
    (24, "An X-ray is used to ___________________________.",
        [("examine the brain", False), ("monitor heart activity", False), ("detect bone fractures", True), ("examine the stomach", False)]),
    (25, "A sphygmomanometer is used to ___________________________.",
        [("examine the brain", False), ("monitor heart activity", False), ("detect bone fractures", False), ("measure blood pressure", True)]),
    (26, "A stethoscope is used to ___________________________.",
        [("examine the brain", False), ("listen to heart and lung sounds", True), ("detect bone fractures", False), ("measure blood pressure", False)]),
    (27, "A thermometer is used to ___________________________.",
        [("examine the brain", False), ("monitor heart activity", False), ("detect bone fractures", False), ("measure body temperature", True)]),
    (28, "A colonoscopy is used to ___________________________.",
        [("examine the stomach", False), ("examine the colon", True), ("monitor heart activity", False), ("detect bone fractures", False)]),
    (29, "A gastroscopy is used to ___________________________.",
        [("examine the stomach", True), ("examine the colon", False), ("monitor heart activity", False), ("detect bone fractures", False)]),
    (30, "An antibiotic is used to ___________________________.",
        [("kill bacteria", True), ("reduce fever", False), ("lower blood pressure", False), ("relieve pain", False)]),
    (31, "An analgesic is used to ___________________________.",
        [("kill bacteria", False), ("reduce fever", False), ("lower blood pressure", False), ("relieve pain", True)]),
    (32, "An antipyretic is used to ___________________________.",
        [("kill bacteria", False), ("reduce fever", True), ("lower blood pressure", False), ("relieve pain", False)]),
    (33, "An antihypertensive is used to ___________________________.",
        [("kill bacteria", False), ("reduce fever", False), ("lower blood pressure", True), ("relieve pain", False)]),
    (34, "A vaccine is used to ___________________________.",
        [("prevent diseases by building immunity", True), ("reduce fever", False), ("lower blood pressure", False), ("relieve pain", False)]),
    (35, "An inhaler is used to ___________________________.",
        [("kill bacteria", False), ("reduce fever", False), ("lower blood pressure", False), ("treat asthma", True)]),
    (36, "An ointment is applied to ___________________________.",
        [("the mouth", False), ("the skin", True), ("the rectum", False), ("the bloodstream", False)]),
    (37, "A suppository is inserted into ___________________________.",
        [("the mouth", False), ("the skin", False), ("the rectum", True), ("the bloodstream", False)]),
    (38, "An injection is administered into ___________________________.",
        [("the mouth", False), ("the skin", False), ("the rectum", False), ("the bloodstream", True)]),
    (39, "Rospotrebnadzor is responsible for ___________________________.",
        [("space exploration", False), ("consumer protection", True), ("transportation", False), ("education", False)]),
    (40, "The main goal of epidemiology is to ___________________________.",
        [("prevent and control diseases", True), ("treat only", False), ("diagnose only", False), ("ignore diseases", False)]),
    (41, "The main purpose of consumer protection is to ___________________________.",
        [("increase prices", False), ("protect buyers from unfair practices", True), ("promote advertising", False), ("reduce product quality", False)]),
    (42, "Hand hygiene is the most effective way to ___________________________.",
        [("treat diseases", False), ("prevent healthcare-associated infections", True), ("diagnose conditions", False), ("perform surgery", False)]),
    (43, "The main source of water pollution in cities is ___________________________.",
        [("parks", False), ("industrial waste", True), ("playgrounds", False), ("schools", False)]),
    (44, "The main cause of foodborne illness is ___________________________.",
        [("chemicals", False), ("bacteria", True), ("radiation", False), ("noise", False)]),
    (45, "Occupational hygiene focuses on ___________________________.",
        [("workplace safety and health", True), ("food safety", False), ("school hygiene", False), ("radiation control", False)]),
    (46, "Community hygiene studies the impact of ___________________________ on health.",
        [("education", False), ("environment", True), ("economy", False), ("politics", False)]),
    (47, "The main function of white blood cells is to ___________________________.",
        [("carry oxygen", False), ("fight infection", True), ("clot blood", False), ("transport nutrients", False)]),
    (48, "A psychiatrist is a doctor who ___________________________.",
        [("treats physical injuries", False), ("performs operations", False), ("diagnoses and treats mental disorders", True), ("specializes in children's diseases", False)]),
    (49, "The main function of platelets is to ___________________________.",
        [("carry oxygen", False), ("fight infection", False), ("help blood clot", True), ("transport nutrients", False)]),
    (50, "The main function of the heart is to ___________________________.",
        [("digest food", False), ("pump blood", True), ("filter waste", False), ("produce hormones", False)]),
    (51, "The main function of the kidneys is to ___________________________.",
        [("pump blood", False), ("digest food", False), ("filter waste from blood", True), ("produce hormones", False)]),
    (52, "The main function of the liver is to ___________________________.",
        [("pump blood", False), ("detoxify the body", True), ("filter air", False), ("digest food", False)]),
    (53, "An endocrinologist treats disorders of ___________________________.",
        [("the heart", False), ("the nervous system", False), ("the endocrine (hormone) system", True), ("the digestive system", False)]),
    (54, "Psychiatry is the branch of medicine that studies ___________________________.",
        [("the heart", False), ("the skin", False), ("mental disorders", True), ("the digestive system", False)]),
    (55, "The best way to prevent the spread of COVID-19 is ___________________________.",
        [("ignoring symptoms", False), ("vaccination and hand hygiene", True), ("sharing food", False), ("going to crowded places", False)]),
    (56, "A pandemic is ___________________________.",
        [("a local outbreak", False), ("a global outbreak of a disease", True), ("a seasonal flu", False), ("a common cold", False)]),
    (57, "An endemic disease is ___________________________.",
        [("constantly present in a certain population", True), ("a global outbreak", False), ("a single case", False), ("a chronic condition", False)]),
    (58, "Isolation is used to ___________________________.",
        [("separate sick individuals from healthy ones", True), ("treat patients", False), ("test vaccines", False), ("close hospitals", False)]),
    (59, "Sterilization is the process of ___________________________.",
        [("cleaning surfaces", False), ("destroying all microorganisms", True), ("reducing germs", False), ("washing hands", False)]),
    (60, "Disinfection is the process of ___________________________.",
        [("destroying all microorganisms", False), ("cleaning surfaces", False), ("reducing germs to a safe level", True), ("washing hands", False)]),
    (61, "Antibiotic resistance occurs when ___________________________.",
        [("bacteria stop responding to antibiotics", True), ("viruses stop responding to antibiotics", False), ("patients stop taking medication", False), ("doctors prescribe too many drugs", False)]),
    (62, "A chronic disease is ___________________________.",
        [("a short-term illness", False), ("a long-lasting condition", True), ("a viral infection", False), ("a foodborne illness", False)]),
    (63, "An acute disease is ___________________________.",
        [("a short-term illness", True), ("a long-lasting condition", False), ("a genetic disorder", False), ("a chronic condition", False)]),
    (64, "The main cause of respiratory infections is ___________________________.",
        [("poor diet", False), ("viruses and bacteria", True), ("lack of sleep", False), ("exercise", False)]),
    (65, "The most common cause of dental disease is ___________________________.",
        [("genetics only", False), ("poor oral hygiene", True), ("drinking water", False), ("eating vegetables", False)]),
    (66, "The doctor ___________ the patient's blood pressure every morning.",
        [("measure", False), ("measures", True), ("measured", False), ("have measured", False)]),
    (67, "The nurse ___________ the patient's temperature yesterday.",
        [("take", False), ("takes", False), ("took", True), ("have taken", False)]),
    (68, "The patient ___________ to the operating room right now.",
        [("is taken", False), ("was taken", False), ("is being taken", True), ("has been taken", False)]),
    (69, "The doctor ___________ already ___________ the results of the test.",
        [("has / checked", True), ("have / checked", False), ("is / checking", False), ("will / check", False)]),
    (70, "The surgeon ___________ the operation at the moment.",
        [("performs", False), ("performed", False), ("is performing", True), ("has performed", False)]),
    (71, "The patient ___________ to the clinic yesterday.",
        [("go", False), ("goes", False), ("went", True), ("have gone", False)]),
    (72, "The ambulance ___________ to the hospital in ten minutes.",
        [("arrive", False), ("arrived", False), ("has arrived", False), ("will arrive", True)]),
    (73, "The nurse ___________ already ___________ the wound.",
        [("has / dressed", True), ("have / dressed", False), ("is / dressing", False), ("will / dress", False)]),
    (74, "The patient ___________ the medication every day.",
        [("take", False), ("takes", True), ("took", False), ("have taken", False)]),
    (75, "The doctor ___________ the patient carefully yesterday.",
        [("examine", False), ("examines", False), ("examined", True), ("have examined", False)]),
    (76, "The patient ___________ better soon.",
        [("feel", False), ("feels", False), ("felt", False), ("will feel", True)]),
    (77, "The nurse ___________ the patient's blood pressure twice a day.",
        [("check", False), ("checks", True), ("checked", False), ("have checked", False)]),
    (78, "The ambulance ___________ at the hospital at 7 pm yesterday.",
        [("arrive", False), ("arrives", False), ("arrived", True), ("have arrived", False)]),
    (79, "The doctor ___________ the patient next week.",
        [("operates", False), ("operated", False), ("has operated", False), ("will operate", True)]),
    (80, "The most effective way to prevent food contamination is ___________.",
        [("using raw ingredients", False), ("proper cooking and storage", True), ("ignoring hygiene", False)]),
    (81, "___________ is used to examine the brain.",
        [("ECG", False), ("X-ray", False), ("MRI", True)]),
    (82, "The main cause of skin diseases is ___________.",
        [("genetics", False), ("poor hygiene", True), ("exercise", False)]),
    (83, "___________ is a type of medication used to reduce fever.",
        [("Antipyretic", True), ("Antibiotic", False), ("Antihypertensive", False)]),
    (84, "The most common cause of radiation exposure is ___________.",
        [("mobile phones", False), ("medical procedures", True), ("cooking", False)]),
    (85, "___________ is a way to reduce pollution in cities.",
        [("Planting trees", True), ("Increasing car usage", False), ("Building more factories", False)]),
    (86, "The main function of a vaccine is to ___________.",
        [("treat diseases", False), ("prevent diseases by building immunity", True), ("cure infections", False)]),
    (87, "___________ is a branch of occupational hygiene.",
        [("Food safety", False), ("Ergonomics", True), ("Radiation hygiene", False)]),
    (88, "The most important measure to prevent the spread of COVID-19 is ___________.",
        [("ignoring symptoms", False), ("vaccination and hand hygiene", True), ("sharing food", False)]),
    (89, "___________ is a long-term health effect of air pollution.",
        [("Headache", False), ("Respiratory diseases", True), ("Skin rash", False)]),
    (90, "The most common preventable risk factor for chronic diseases is ___________.",
        [("genetics", False), ("smoking", True), ("physical exercise", False)]),
]

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.font.bold = False
rpr = style.element.get_or_add_rPr()
rfonts = rpr.get_or_add_rFonts()
rfonts.set(qn('w:ascii'), 'Times New Roman')
rfonts.set(qn('w:hAnsi'), 'Times New Roman')
rfonts.set(qn('w:cs'), 'Times New Roman')
pf = style.paragraph_format
pf.line_spacing = 1.0
pf.space_before = Pt(0)
pf.space_after = Pt(0)

def set_cell_text(cell, text, align, vcenter=False):
    cell.text = ""
    if vcenter:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = False
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:ascii'), 'Times New Roman')
    rf.set(qn('w:hAnsi'), 'Times New Roman')
    rf.set(qn('w:cs'), 'Times New Roman')

def add_plain(text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = bold
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:ascii'), 'Times New Roman')
    rf.set(qn('w:hAnsi'), 'Times New Roman')
    rf.set(qn('w:cs'), 'Times New Roman')
    return p

def set_col_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn('w:tblLayout'))
    if layout is None:
        layout = tblPr.makeelement(qn('w:tblLayout'), {})
        tblPr.append(layout)
    layout.set(qn('w:type'), 'fixed')
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w

CENTER = WD_ALIGN_PARAGRAPH.CENTER
LEFT = WD_ALIGN_PARAGRAPH.LEFT
COL_WIDTHS = [Cm(2.4), Cm(11.5), Cm(3.5)]

for num, qtext, options in questions:
    # correct answer (+) always first, rest keep original order
    options = sorted(options, key=lambda o: not o[1])
    add_plain("%d. %s" % (num, qtext))

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    hdr = table.rows[0].cells
    set_cell_text(hdr[0], "Поле для выбора ответа", CENTER, vcenter=True)
    set_cell_text(hdr[1], "Варианты ответов", CENTER, vcenter=True)
    set_cell_text(hdr[2], "Поле для отметки правильного ответа (+)", CENTER)

    for opt_text, correct in options:
        row = table.add_row().cells
        set_cell_text(row[0], "", CENTER, vcenter=True)
        set_cell_text(row[1], opt_text, LEFT, vcenter=True)
        set_cell_text(row[2], "+" if correct else "-", LEFT)

    set_col_widths(table, COL_WIDTHS)

doc.save("quiz2.docx")
print("saved quiz2.docx with", len(questions), "questions")
